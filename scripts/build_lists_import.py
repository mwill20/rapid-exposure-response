# Generates lists/RER_Lists_Import.xlsx: one Excel table per SharePoint list, matching lists/DATA_MODEL.md.
# Usage: python scripts/build_lists_import.py <output.xlsx>
# Each sheet has one EXAMPLE row so "Create a list from Excel" infers date and number types. Delete it after import.
import os
import sys
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

F = 'Arial'
HDR = PatternFill('solid', start_color='1F3864')
now = datetime(2026, 10, 20, 18, 0)
DATE_FMT = 'yyyy-mm-dd hh:mm'

wb = Workbook()
wb.remove(wb.active)


def table(name, columns, example, widths=None):
    ws = wb.create_sheet(name)
    for i, h in enumerate(columns, 1):
        c = ws.cell(1, i, h)
        c.font = Font(name=F, bold=True, color='FFFFFF')
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[c.column_letter].width = (widths or {}).get(h, 18)
    for i, h in enumerate(columns, 1):
        v = example.get(h, '')
        c = ws.cell(2, i, v)
        c.font = Font(name=F, italic=True)
        if isinstance(v, datetime):
            c.number_format = DATE_FMT
    last = ws.cell(1, len(columns)).column_letter
    t = Table(displayName=name, ref=f'A1:{last}2')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    ws.add_table(t)
    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 32
    return ws


table('RER_Events',
      ['Event_ID', 'Name', 'CVEs', 'Source', 'Source_Tier', 'Lane', 'Stage', 'Blocker', 'Declared_At', 'Triage_Lead',
       'Owners_Identified_At', 'Exposure_Answered_At', 'First_Control_At', 'Fix_Verified_At', 'Hunt_Complete_At',
       'Hunt_Result', 'Sensitive_Data_Impact', 'Compliance_Notified_At', 'Status', 'Closed_At', 'Event_Pack_Link', 'AAR_Link'],
      {'Event_ID': 'EVT-2026-001', 'Name': 'EXAMPLE ROW: delete after import', 'CVEs': 'EXAMPLE-2026-0001',
       'Source': 'Vendor advisory', 'Source_Tier': 'Tier 1', 'Lane': 'Fire drill', 'Stage': 'S2 Contain',
       'Blocker': 'EXAMPLE: WAF rule not yet verified', 'Declared_At': now, 'Triage_Lead': 'Jane Example',
       'Owners_Identified_At': now + timedelta(hours=1.5), 'Exposure_Answered_At': now + timedelta(hours=5),
       'First_Control_At': now + timedelta(hours=9), 'Fix_Verified_At': now + timedelta(hours=40),
       'Hunt_Complete_At': now + timedelta(hours=46), 'Hunt_Result': 'Pending', 'Sensitive_Data_Impact': 'Unknown',
       'Compliance_Notified_At': now + timedelta(hours=2), 'Status': 'Open', 'Closed_At': now + timedelta(days=5),
       'Event_Pack_Link': 'https://example.org/RER/org/events/EVT-2026-001.docx', 'AAR_Link': 'https://example.org/RER/org/events/EVT-2026-001-AAR.docx'},
      {'Name': 34, 'Blocker': 40, 'Event_Pack_Link': 40, 'AAR_Link': 40})

table('RER_Assets',
      ['Asset_Ref', 'Event_ID', 'Asset_ID', 'System_Name', 'Owner', 'Version_Found', 'Q1_Runs_It', 'Q2_Vuln_Version',
       'Q3_Internet_NoAuth', 'Q4_Internal_Broad', 'Q5_Feature_Enabled', 'Q6_Segmentation_Verified', 'Q7_Sensitive_Data',
       'DT1_Row', 'Grade', 'Priority', 'Owner_Notified_At', 'Owner_Ack_At', 'Remediation_Status', 'Verified_At',
       'Evidence_Link', 'Notes'],
      {'Asset_Ref': 'EVT-2026-001 / AST-001', 'Event_ID': 'EVT-2026-001', 'Asset_ID': 'AST-001',
       'System_Name': 'EXAMPLE web tier', 'Owner': 'Jane Example', 'Version_Found': '7.2.1', 'Q1_Runs_It': 'Y',
       'Q2_Vuln_Version': 'Y', 'Q3_Internet_NoAuth': 'Y', 'Q4_Internal_Broad': 'Y', 'Q5_Feature_Enabled': 'Y',
       'Q6_Segmentation_Verified': 'N', 'Q7_Sensitive_Data': 'Y', 'DT1_Row': 4, 'Grade': 'E1', 'Priority': 'P1',
       'Owner_Notified_At': now + timedelta(hours=1.5), 'Owner_Ack_At': now + timedelta(hours=2),
       'Remediation_Status': 'Control in place', 'Verified_At': now + timedelta(hours=40),
       'Evidence_Link': 'https://example.org/RER/org/evidence/EV-001', 'Notes': 'EXAMPLE ROW: delete after import'},
      {'Asset_Ref': 24, 'System_Name': 26, 'Evidence_Link': 40, 'Notes': 34})

table('RER_Controls',
      ['Log_ID', 'Event_ID', 'Asset_ID', 'Control_ID', 'DT2_Row', 'Determination', 'Implemented_At', 'Implemented_By',
       'Approved_By', 'Rollback_Plan', 'Expiry_Date', 'Status', 'Verified', 'Evidence_Link'],
      {'Log_ID': 'TC-001', 'Event_ID': 'EVT-2026-001', 'Asset_ID': 'AST-001', 'Control_ID': 'CTL-03', 'DT2_Row': 3,
       'Determination': 'EXAMPLE: Restricting the management interface removes it from reach of the internet. If the flaw is on the user-facing side this does not cover it, so the patch stays urgent.',
       'Implemented_At': now + timedelta(hours=9), 'Implemented_By': 'Network team', 'Approved_By': 'Jane Example',
       'Rollback_Plan': 'Remove ACL', 'Expiry_Date': now + timedelta(days=30), 'Status': 'Active', 'Verified': 'Y',
       'Evidence_Link': 'https://example.org/RER/org/evidence/EV-001'},
      {'Determination': 60, 'Rollback_Plan': 24, 'Evidence_Link': 40})

table('RER_Hunts',
      ['Hunt_ID', 'Event_ID', 'Hypothesis', 'Required_Data', 'Data_Present_Full_Window', 'Window_Start', 'Window_End',
       'Outcome', 'Run_By', 'Run_At', 'Query_Link', 'Notes'],
      {'Hunt_ID': 'HT-001', 'Event_ID': 'EVT-2026-001', 'Hypothesis': 'H7',
       'Required_Data': 'Device admin and audit logs, identity provider sign-in logs', 'Data_Present_Full_Window': 'N',
       'Window_Start': now - timedelta(days=21), 'Window_End': now, 'Outcome': 'Insufficient data', 'Run_By': 'Jane Example',
       'Run_At': now + timedelta(hours=30), 'Query_Link': 'https://example.org/siem/saved-search/EXAMPLE',
       'Notes': 'EXAMPLE ROW: device audit log retention is 14 days; gap logged for the AAR'},
      {'Required_Data': 40, 'Query_Link': 40, 'Notes': 44})

table('RER_Evidence',
      ['Evidence_ID', 'Event_ID', 'Asset_ID', 'Item', 'Standard_Met', 'Captured_By', 'Captured_At', 'Link', 'Notes'],
      {'Evidence_ID': 'EV-001', 'Event_ID': 'EVT-2026-001', 'Asset_ID': 'AST-001', 'Item': 'CTL-03', 'Standard_Met': 'Y',
       'Captured_By': 'Jane Example', 'Captured_At': now + timedelta(hours=10),
       'Link': 'https://example.org/RER/org/evidence/EV-001',
       'Notes': 'EXAMPLE ROW: rule export plus allowed and disallowed vantage tests, timestamps and hostname visible'},
      {'Link': 40, 'Notes': 50})

out = sys.argv[1] if len(sys.argv) > 1 else 'RER_Lists_Import.xlsx'
os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
wb.save(out)
print('saved', out)
