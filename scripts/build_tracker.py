# Generates the RER Tracker workbook. Usage: python scripts/build_tracker.py <output.xlsx>
# Excel recalculates formulas on open.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime, timedelta

N = 200  # data rows per sheet
F = 'Arial'
HDR = PatternFill('solid', start_color='1F3864')
FILL = PatternFill('solid', start_color='FFF2CC')   # user input
FORM = PatternFill('solid', start_color='EDEDED')   # formula
thin = Side(style='thin', color='BFBFBF')
wb = Workbook()

def sheet(name, cols, widths, formulas=None, example=None, dvs=None, date_cols=()):
    ws = wb.create_sheet(name)
    formulas = formulas or {}
    for i, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(1, i, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[c.column_letter].width = w
    for r in range(2, N + 2):
        for i, h in enumerate(cols, 1):
            c = ws.cell(r, i)
            c.font = Font(name=F, italic=(r == 2))
            if h in formulas:
                c.value = formulas[h].replace('{r}', str(r)); c.fill = FORM
            else:
                c.fill = FILL
            if h in date_cols: c.number_format = 'yyyy-mm-dd hh:mm'
    if example:
        for i, h in enumerate(cols, 1):
            if h not in formulas and h in example: ws.cell(2, i, example[h])
    for col_letter, options in (dvs or {}).items():
        dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f'{col_letter}2:{col_letter}{N+1}')
    ws.freeze_panes = 'B2'
    last = ws.cell(1, len(cols)).column_letter
    t = Table(displayName='t' + name.replace('_', ''), ref=f'A1:{last}{N+1}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    ws.add_table(t)
    ws.row_dimensions[1].height = 32
    return ws

# README
rd = wb.active; rd.title = 'README'
lines = [
 ('RER Tracker: Rapid Exposure Response', True),
 ('Companion to the RER Playbook. Store both in the same restricted SharePoint location.', False),
 ('', False),
 ('LEGEND', True),
 ('Yellow cells: you fill these in.', False),
 ('Gray cells: formulas. Do not overwrite.', False),
 ('Row 2 on every sheet is an EXAMPLE (italic). Overwrite or delete it once you add real data.', False),
 ('Dates: enter as yyyy-mm-dd hh:mm so the hour calculations work.', False),
 ('', False),
 ('SHEETS', True),
 ('Registry: your Top-25 (then growing) list of systems, components, exposure, and owner. Start with internet-facing systems.', False),
 ('Owners: one row per person. Registry and Event_Assets look names up from here by Owner_ID.', False),
 ('Controls_Menu: pre-approved compensating controls (matches Playbook Section 7).', False),
 ('Events: one row per RER event. Hour columns calculate your response times automatically.', False),
 ('Event_Assets: one row per affected asset per event. Drives owner notices and verification.', False),
 ('Temp_Controls: every temporary control with expiry. Days_Left and Flag show what is overdue.', False),
 ('Metrics: summary for leadership updates and after-action reviews.', False),
 ('', False),
 ('COPILOT TIPS', True),
 ('Add this workbook as a reference in the RER Copilot Notebook or agent. Then ask, for example:', False),
 ('"Which Temp_Controls rows are Overdue?"  "List Event_Assets for EVT-2026-001 that are not Verified."', False),
 ('In Excel, Copilot can also add columns or summarize a sheet. Review any formula it writes before relying on it.', False),
 ('', False),
 ('SECURITY', True),
 ('This workbook is an attack-surface map. Classify it as [FILL: classification] and limit access to RER roles.', False),
 ('Copilot surfaces content to anyone who has access, so check the site and file permissions before sharing.', False),
 ('Never store credentials, secrets, or patient data here.', False),
]
for i, (t, b) in enumerate(lines, 1):
    c = rd.cell(i, 1, t); c.font = Font(name=F, bold=b, size=14 if i == 1 else 11, color='1F3864' if b else '000000')
rd.column_dimensions['A'].width = 120

now = datetime(2026, 10, 20, 18, 0)

sheet('Owners',
 ['Owner_ID','Name','Role_Team','Email','Phone_or_Teams','Manager','After_Hours_Reachable','Notes'],
 [11,22,22,28,18,20,14,30],
 example={'Owner_ID':'OWN-001','Name':'Jane Example','Role_Team':'Web Platform','Email':'jane@example.org','Phone_or_Teams':'Teams','Manager':'Sam Example','After_Hours_Reachable':'Y','Notes':'EXAMPLE ROW'},
 dvs={'G':['Y','N']})

sheet('Registry',
 ['Asset_ID','System_Name','Component','Version','Hosting','Exposure','Data_Sensitivity','Criticality_1to5','Owner_ID','Owner_Name','Backup_Owner_ID','Vendor_Advisory_URL','Last_Verified','Notes'],
 [10,24,18,10,15,15,14,11,10,20,11,30,14,28],
 formulas={'Owner_Name':'=IF(I{r}="","",IFERROR(INDEX(Owners!$B:$B,MATCH(I{r},Owners!$A:$A,0)),"NOT FOUND"))'},
 example={'Asset_ID':'AST-001','System_Name':'Patient portal web tier','Component':'PHP','Version':'8.3.30','Hosting':'On-prem','Exposure':'Internet','Data_Sensitivity':'ePHI','Criticality_1to5':5,'Owner_ID':'OWN-001','Backup_Owner_ID':'','Vendor_Advisory_URL':'https://www.php.net/','Last_Verified':now,'Notes':'EXAMPLE ROW'},
 dvs={'E':['On-prem','Cloud','SaaS','Vendor-managed'],'F':['Internet','Internal-broad','Segmented','Unknown'],'G':['ePHI','PII','Internal','Public'],'H':['1','2','3','4','5']},
 date_cols=('Last_Verified',))

ctrl = sheet('Controls_Menu',
 ['Control_ID','Control','Use_When','Tier','Approver','Executed_By','Rollback','Max_Days'],
 [11,32,40,7,20,20,20,10])
menu = [
 ('CTL-01','WAF / edge virtual patch','Trusted source publishes request indicators','A','Triage Lead','[FILL]','Disable rule',30),
 ('CTL-02','Disable vulnerable feature or module','Vendor names a feature-level mitigation','B','[FILL]','[FILL]','Revert config',30),
 ('CTL-03','Restrict access (allowlist, VPN-only)','Service does not need broad reach','B','[FILL]','[FILL]','Remove ACL',30),
 ('CTL-04','Segment or isolate','Patch far off and exposure high','B','[FILL]','[FILL]','Restore path',14),
 ('CTL-05','Extra logging plus detection rule','Always; mandatory when mitigations are partial','A','Triage Lead','[FILL]','N/A',30),
 ('CTL-06','Credential, session, or key rotation','Possible compromise or exposed secrets','B','[FILL]','[FILL]','N/A',1),
 ('CTL-07','Take offline','E1, exploited, no other option','C','[FILL]','[FILL]','Restore service',7),
]
for r, row in enumerate(menu, 2):
    for i, v in enumerate(row, 1):
        c = ctrl.cell(r, i, v); c.font = Font(name=F)

hrs = lambda a, b: f'=IF(OR({a}{{r}}="",{b}{{r}}=""),"",ROUND(({b}{{r}}-{a}{{r}})*24,1))'
sheet('Events',
 ['Event_ID','Name','CVEs','Source','Source_Tier','Lane','Declared_At','Triage_Lead','Owners_Identified_At','Exposure_Answered_At','First_Control_At','Fix_Verified_At','Hunt_Result','Sensitive_Data_Impact','Compliance_Notified_At','Status','Closed_At','Hrs_to_Owners','Hrs_to_Exposure','Hrs_to_Control','Hrs_to_Verified','AAR_Link'],
 [14,22,26,22,10,11,16,16,16,16,16,16,14,12,16,11,16,10,10,10,10,24],
 formulas={'Hrs_to_Owners':hrs('G','I'),'Hrs_to_Exposure':hrs('G','J'),'Hrs_to_Control':hrs('G','K'),'Hrs_to_Verified':hrs('G','L')},
 example={'Event_ID':'EVT-2026-001','Name':'PHP2Shell (EXAMPLE)','CVEs':'CVE-2026-109618, CVE-2026-109619','Source':'PHP project advisory','Source_Tier':'Tier 1','Lane':'Fire drill','Declared_At':now,'Triage_Lead':'Jane Example','Owners_Identified_At':now+timedelta(hours=1.5),'Exposure_Answered_At':now+timedelta(hours=5),'First_Control_At':now+timedelta(hours=9),'Fix_Verified_At':now+timedelta(hours=40),'Hunt_Result':'No evidence','Sensitive_Data_Impact':'Unknown','Status':'Open'},
 dvs={'E':['Tier 1','Tier 2','Tier 3'],'F':['Prepare','Fire drill','Routine'],'M':['Pending','No evidence','Suspicious','Confirmed'],'N':['Y','N','Unknown'],'P':['Open','Monitoring','Closed']},
 date_cols=('Declared_At','Owners_Identified_At','Exposure_Answered_At','First_Control_At','Fix_Verified_At','Compliance_Notified_At','Closed_At'))

sheet('Event_Assets',
 ['Event_ID','Asset_ID','System_Name','Owner_Name','Version_Found','Vulnerable','Exposure_Grade','Owner_Notified_At','Owner_Ack_At','Remediation_Status','Verified_By','Evidence_Link','Notes'],
 [14,10,24,20,12,11,11,16,16,18,16,26,26],
 formulas={'System_Name':'=IF(B{r}="","",IFERROR(INDEX(Registry!$B:$B,MATCH(B{r},Registry!$A:$A,0)),"NOT IN REGISTRY"))',
           'Owner_Name':'=IF(B{r}="","",IFERROR(INDEX(Registry!$J:$J,MATCH(B{r},Registry!$A:$A,0)),""))'},
 example={'Event_ID':'EVT-2026-001','Asset_ID':'AST-001','Version_Found':'8.3.30','Vulnerable':'Y','Exposure_Grade':'E1','Owner_Notified_At':now+timedelta(hours=1.5),'Owner_Ack_At':now+timedelta(hours=2),'Remediation_Status':'Verified','Verified_By':'Jane Example','Notes':'EXAMPLE ROW'},
 dvs={'F':['Y','N','Unknown'],'G':['E1','E2','E3','E4'],'J':['Not started','Control in place','Patched','Verified','Not affected']},
 date_cols=('Owner_Notified_At','Owner_Ack_At'))

sheet('Temp_Controls',
 ['Log_ID','Event_ID','Asset_ID','Control_ID','Control_Name','Implemented_At','Implemented_By','Approved_By','Rollback_Plan','Expiry_Date','Status','Days_Left','Flag'],
 [10,14,10,11,30,16,16,16,22,14,11,10,12],
 formulas={'Control_Name':'=IF(D{r}="","",IFERROR(INDEX(Controls_Menu!$B:$B,MATCH(D{r},Controls_Menu!$A:$A,0)),"UNKNOWN ID"))',
           'Days_Left':'=IF(OR(J{r}="",K{r}="Removed"),"",ROUND(J{r}-TODAY(),0))',
           'Flag':'=IF(L{r}="","",IF(L{r}<0,"OVERDUE",IF(L{r}<=3,"Due soon","OK")))'},
 example={'Log_ID':'TC-001','Event_ID':'EVT-2026-001','Asset_ID':'AST-001','Control_ID':'CTL-01','Implemented_At':now+timedelta(hours=9),'Implemented_By':'WAF team','Approved_By':'Jane Example','Rollback_Plan':'Disable rule in WAF console','Expiry_Date':datetime(2026,11,20),'Status':'Active'},
 dvs={'D':[m[0] for m in menu],'K':['Active','Extended','Removed']},
 date_cols=('Implemented_At',))
wb['Temp_Controls'].column_dimensions['J'].width = 14
for r in range(2, N+2): wb['Temp_Controls'].cell(r, 10).number_format = 'yyyy-mm-dd'

m = wb.create_sheet('Metrics')
m.column_dimensions['A'].width = 46; m.column_dimensions['B'].width = 16; m.column_dimensions['C'].width = 50
rng = lambda s, col: f"{s}!${col}$2:${col}${N+1}"
rows = [
 ('Metric','Value','How it is calculated'),
 ('Open events', f'=COUNTIF({rng("Events","P")},"Open")', 'Events with Status = Open'),
 ('Avg hours to owners identified', f'=IFERROR(ROUND(AVERAGE({rng("Events","R")}),1),"n/a")', 'Average of Events Hrs_to_Owners'),
 ('Avg hours to exposure answered', f'=IFERROR(ROUND(AVERAGE({rng("Events","S")}),1),"n/a")', 'Average of Events Hrs_to_Exposure'),
 ('Avg hours to first control', f'=IFERROR(ROUND(AVERAGE({rng("Events","T")}),1),"n/a")', 'Average of Events Hrs_to_Control'),
 ('Avg hours to fix verified', f'=IFERROR(ROUND(AVERAGE({rng("Events","U")}),1),"n/a")', 'Average of Events Hrs_to_Verified'),
 ('E1 assets not yet Verified', f'=COUNTIFS({rng("Event_Assets","G")},"E1",{rng("Event_Assets","J")},"<>Verified",{rng("Event_Assets","J")},"<>Not affected")', 'Event_Assets with grade E1 and status not Verified / Not affected'),
 ('Owners not yet acknowledged', f'=COUNTIFS({rng("Event_Assets","H")},"<>",{rng("Event_Assets","I")},"")', 'Notified but no Owner_Ack_At'),
 ('Temporary controls OVERDUE', f'=COUNTIF({rng("Temp_Controls","M")},"OVERDUE")', 'Temp_Controls Flag = OVERDUE'),
 ('Temporary controls due within 3 days', f'=COUNTIF({rng("Temp_Controls","M")},"Due soon")', 'Temp_Controls Flag = Due soon'),
 ('Orphan assets (no owner)', f'=COUNTIFS({rng("Registry","A")},"<>",{rng("Registry","I")},"")', 'Registry rows with Asset_ID but no Owner_ID'),
 ('Owner IDs not found in Owners', f'=COUNTIF({rng("Registry","J")},"NOT FOUND")', 'Registry Owner_ID that does not match Owners'),
]
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row, 1):
        c = m.cell(i, j, v); c.font = Font(name=F, bold=(i == 1), color='FFFFFF' if i == 1 else '000000')
        if i == 1: c.fill = HDR
        elif j == 2: c.fill = FORM
m.cell(len(rows)+2, 1, 'Note: averages include the EXAMPLE row until you delete it.').font = Font(name=F, italic=True, color='777777')

import sys, os
out = sys.argv[1] if len(sys.argv) > 1 else 'RER_Tracker.xlsx'
os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
wb.save(out); print('saved', out)
